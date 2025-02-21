import tkinter as tk
from tkinter import ttk, filedialog
import os
import yaml
from typing import Dict, List
from collections import defaultdict
from openpyxl import load_workbook

class YamlClassifierApp:
    def __init__(self, root):
        self.root = root
        self.root.title("YAML Classifier")
        self.root.geometry("1200x600")
        
        # 座位位置映射
        self.seat_mapping = {
            'seatleftrow1': 'FL',
            'seatleftrow2': 'RL',
            'seatrightrow1': 'FR',
            'seatrightrow2': 'RR',
            'seatmidrow2': 'RM'
        }
        
        # 配置根窗口的网格权重
        self.root.grid_rowconfigure(0, weight=1)
        self.root.grid_columnconfigure(0, weight=1)
        
        # 存储组合分类结果
        self.combined_results = defaultdict(list)
        
        # 定义分类顺序
        self.classification_order = {
            'handle': 0,    # handle_folded 或 handle_upright
            'freq': 1,      # breathing frequency
            'mode': 2,      # breathing mode
            'seat': 3       # seat_position
        }
        
        self.setup_ui()
        
    def setup_ui(self):
        # 创建主框架
        main_frame = ttk.Frame(self.root)
        main_frame.grid(row=0, column=0, sticky='nsew', padx=10, pady=10)
        
        # 配置主框架的网格权重
        main_frame.grid_rowconfigure(1, weight=1)
        main_frame.grid_columnconfigure(0, weight=1)
        
        # 按钮框架
        button_frame = ttk.Frame(main_frame)
        button_frame.grid(row=0, column=0, sticky='ew', pady=(0, 10))
        
        # YAML文件夹选择按钮
        ttk.Button(button_frame, text="选择YAML文件夹", command=self.select_folder).pack(side='left', padx=5)
        
        # Excel文件选择按钮
        ttk.Button(button_frame, text="选择Excel文件", command=self.select_excel).pack(side='left', padx=5)
        
        # Excel文件路径显示
        self.excel_path_var = tk.StringVar()
        ttk.Label(button_frame, textvariable=self.excel_path_var).pack(side='left', padx=5)
        
        # 搜索框架
        search_frame = ttk.Frame(main_frame)
        search_frame.grid(row=1, column=0, sticky='ew', pady=(0, 10))
        
        # 搜索标签
        ttk.Label(search_frame, text="搜索:").pack(side='left', padx=5)
        
        # 搜索输入框
        self.search_var = tk.StringVar()
        self.search_entry = ttk.Entry(search_frame, textvariable=self.search_var)
        self.search_entry.pack(side='left', fill='x', expand=True, padx=5)
        
        # 搜索按钮
        ttk.Button(search_frame, text="搜索", command=self.search_results).pack(side='left', padx=5)
        
        # 清除搜索按钮
        ttk.Button(search_frame, text="清除搜索", command=self.clear_search).pack(side='left', padx=5)
        
        # 创建结果显示区域
        result_frame = ttk.LabelFrame(main_frame, text="分类结果")
        result_frame.grid(row=2, column=0, sticky='nsew')
        result_frame.grid_rowconfigure(0, weight=1)
        result_frame.grid_columnconfigure(0, weight=1)
        
        # 创建文本框和滚动条的容器
        text_container = ttk.Frame(result_frame)
        text_container.grid(row=0, column=0, sticky='nsew', padx=5, pady=5)
        text_container.grid_rowconfigure(0, weight=1)
        text_container.grid_columnconfigure(0, weight=1)
        
        # 结果文本框
        self.result_text = tk.Text(text_container, wrap='none')
        self.result_text.grid(row=0, column=0, sticky='nsew')
        
        # 垂直滚动条
        v_scrollbar = ttk.Scrollbar(text_container, orient="vertical", command=self.result_text.yview)
        v_scrollbar.grid(row=0, column=1, sticky='ns')
        
        # 水平滚动条
        h_scrollbar = ttk.Scrollbar(text_container, orient="horizontal", command=self.result_text.xview)
        h_scrollbar.grid(row=1, column=0, sticky='ew')
        
        # 配置文本框的滚动
        self.result_text.configure(
            yscrollcommand=v_scrollbar.set,
            xscrollcommand=h_scrollbar.set,
            font=('Courier', 10)  # 使用等宽字体
        )
        
        # 底部按钮和状态栏容器
        bottom_frame = ttk.Frame(main_frame)
        bottom_frame.grid(row=3, column=0, sticky='ew', pady=(10, 0))
        
        # 复制按钮
        ttk.Button(bottom_frame, text="复制全部结果", command=self.copy_results).pack(side='left')
        
        # 状态栏
        self.status_var = tk.StringVar()
        ttk.Label(bottom_frame, textvariable=self.status_var).pack(side='right')
        
        # 绑定回车键到搜索功能
        self.search_entry.bind('<Return>', lambda e: self.search_results())
        
        # 存储原始结果
        self.original_results = ""

    def select_folder(self):
        folder = filedialog.askdirectory()
        if folder:
            self.process_folder(folder)
    
    def process_folder(self, folder_path):
        # 清空之前的结果
        self.combined_results.clear()
        self.result_text.delete('1.0', tk.END)
        
        # 扫描文件夹中的所有yaml文件
        yaml_files = [f for f in os.listdir(folder_path) if f.endswith('.yaml')]
        processed_files = 0
        
        # 添加调试信息
        print(f"找到的yaml文件: {yaml_files}")
        self.result_text.insert(tk.END, f"正在处理文件夹: {folder_path}\n")
        self.result_text.insert(tk.END, f"找到 {len(yaml_files)} 个yaml文件\n\n")
        
        for file in yaml_files:
            file_path = os.path.join(folder_path, file)
            try:
                with open(file_path, 'r', encoding='utf-8') as f:
                    yaml_data = yaml.safe_load(f)
                    if yaml_data and 'TestCase' in yaml_data:
                        print(f"正在处理文件: {file}")
                        self.classify_yaml(file, yaml_data)
                        processed_files += 1
                    else:
                        print(f"文件格式不正确: {file}")
            except Exception as e:
                error_msg = f"处理文件 {file} 时出错: {str(e)}"
                print(error_msg)
                self.result_text.insert(tk.END, error_msg + "\n")
        
        # 更新显示
        self.update_display()
        
        # 如果已经选择了Excel文件，就更新Excel
        if hasattr(self, 'excel_file_path'):
            print("正在更新Excel文件...")
            self.update_excel()
        
        # 更新状态栏
        status_msg = f"处理完成，共扫描了 {len(yaml_files)} 个文件，成功处理 {processed_files} 个"
        self.status_var.set(status_msg)
        print(status_msg)
        print("分类结果:", self.combined_results)
    
    def classify_yaml(self, file_name: str, yaml_data: dict):
        # 存储当前文件的分类结果
        current_classifications = []
        
        # 检查Vehicle部分
        if 'TestCase' in yaml_data and 'Vehicle' in yaml_data['TestCase']:
            vehicle = yaml_data['TestCase']['Vehicle']
            for key, value in vehicle.items():
                if key.startswith('Seat') and isinstance(value, dict):
                    # 检查是否有Dummy
                    occupancy_seat = value.get('OccupancySeat', '')
                    if occupancy_seat and 'Dummy_' in str(occupancy_seat):
                        # 添加Seat位置
                        seat_key = key.lower().replace('seat', 'seat_')
                        current_classifications.append(seat_key)
                        print(f"找到座位: {seat_key}")
                        
                        # 添加Comment类型
                        comment = str(value.get('OccupancyComment', '')).lower()
                        if 'handle folded back' in comment or 'folded' in comment:
                            current_classifications.append('handle_folded')
                            print("找到handle_folded")
                        elif 'handle upright' in comment or 'upright' in comment:
                            current_classifications.append('handle_upright')
                            print("找到handle_upright")
                        
                        # 检查ChildSeat信息
                        child_seat_model = value.get('ChildSeatModel', '')
                        child_seat_orientation = value.get('ChildSeatOrientation', '')
                        if child_seat_model and child_seat_orientation:
                            print(f"找到儿童座椅: {child_seat_model}, 方向: {child_seat_orientation}")
        
        # 检查Catalog中的Dummy信息
        if 'Catalog' in yaml_data and isinstance(yaml_data['Catalog'], dict):
            catalog = yaml_data['Catalog']
            if 'Dummy' in catalog and isinstance(catalog['Dummy'], list) and len(catalog['Dummy']) > 0:
                # 只处理第一个Dummy的信息
                dummy = catalog['Dummy'][0]
                
                # 添加BreathingMode
                mode = str(dummy.get('BreathingMode', '')).lower()
                if mode:  # 确保mode不为空
                    # 处理特殊的模式名称
                    mode_mapping = {
                        'deep sleep': 'deepsleep',
                        'wake up': 'wakeup',
                        'sleep': 'sleep',
                        'awake': 'awake'
                    }
                    mode = mode_mapping.get(mode, mode)
                    mode_key = f"mode_{mode}"
                    current_classifications.append(mode_key)
                    print(f"找到mode: {mode_key}")
                
                # 添加BreathingFrequency
                freq = dummy.get('BreathingFrequency')
                if freq is not None:
                    freq_key = f"freq_{freq}"
                    current_classifications.append(freq_key)
                    print(f"找到frequency: {freq_key}")
            else:
                print("Catalog.Dummy是空列表或不是列表格式")
        
        print(f"当前文件的分类结果: {current_classifications}")
        
        # 检查是否包含所有必要的分类类型
        has_handle = any(c.startswith('handle_') for c in current_classifications)
        has_seat = any(c.startswith('seat_') for c in current_classifications)
        has_mode = any(c.startswith('mode_') for c in current_classifications)
        has_freq = any(c.startswith('freq_') for c in current_classifications)
        
        if has_handle and has_seat and has_mode and has_freq:
            # 按照指定顺序排序分类
            sorted_classifications = []
            for prefix in ['handle', 'freq', 'mode', 'seat']:
                for item in current_classifications:
                    if item.startswith(prefix):
                        sorted_classifications.append(item)
                        break  # 每种类型只取第一个匹配的
            
            # 组合成key
            key = ','.join(sorted_classifications)
            # 去掉.yaml后缀
            file_name_without_ext = os.path.splitext(file_name)[0]
            self.combined_results[key].append(file_name_without_ext)
            print(f"添加到结果: {key} -> {file_name_without_ext}")
        else:
            missing = []
            if not has_handle: missing.append('handle')
            if not has_seat: missing.append('seat')
            if not has_mode: missing.append('mode')
            if not has_freq: missing.append('freq')
            print(f"文件 {file_name} 缺少以下分类: {', '.join(missing)}")
    
    def update_display(self):
        # 清空显示
        self.result_text.delete('1.0', tk.END)
        self.original_results = ""  # 清空原始结果
        
        # 显示组合结果
        for classifications, files in self.combined_results.items():
            # 格式化输出
            result_line = f"{classifications} : {';'.join(files)}\n\n"
            self.result_text.insert(tk.END, result_line)
    
    def copy_results(self):
        content = self.result_text.get('1.0', 'end-1c')
        self.root.clipboard_clear()
        self.root.clipboard_append(content)
        self.status_var.set("内容已复制到剪贴板")

    def select_excel(self):
        file_path = filedialog.askopenfilename(
            title="选择Excel文件",
            filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")]
        )
        if file_path:
            self.excel_path_var.set(os.path.basename(file_path))
            self.excel_file_path = file_path
            print(f"已选择Excel文件: {file_path}")
            
            # 如果已经有分类结果，立即更新Excel
            if self.combined_results:
                print("检测到已有分类结果，开始更新Excel...")
                self.update_excel()
            else:
                print("还没有YAML分类结果，请先选择YAML文件夹")
                
    def update_excel(self):
        if not hasattr(self, 'excel_file_path'):
            print("请先选择Excel文件")
            return
            
        if not self.combined_results:
            print("没有分类结果可以写入Excel")
            return
            
        try:
            print("开始更新Excel文件...")
            print(f"当前的分类结果: {dict(self.combined_results)}")
            
            # 加载Excel文件
            wb = load_workbook(self.excel_file_path)
            ws = wb.active
            
            # 查找yaml filename列
            yaml_col = None
            for col in range(1, ws.max_column + 1):
                if ws.cell(row=1, column=col).value == 'yaml filename':
                    yaml_col = col
                    break
            
            if yaml_col is None:
                print("未找到'yaml filename'列")
                return
                
            print(f"找到yaml filename列: {yaml_col}")
            
            # 用于存储当前组的共用信息
            current_mode = None
            current_handle = None
            
            # 遍历Excel的每一行
            for row in range(2, ws.max_row + 1):
                # 获取当前行的信息
                seat_pos = ws.cell(row=row, column=3).value  # C列是SeatingPosition1
                breathing_mode = ws.cell(row=row, column=9).value  # I列是Breathing Mode
                occupancy_comment = ws.cell(row=row, column=8).value or ''  # H列是OccupancyComment
                breathing_freq = ws.cell(row=row, column=11).value  # K列是BreathingFrequency
                
                print(f"\n处理第{row}行:")
                print(f"座位: {seat_pos}, 模式: {breathing_mode}, 手柄信息: {occupancy_comment}, 频率: {breathing_freq}")
                
                # 跳过空行或者SUM行
                if not seat_pos or str(seat_pos).strip() == '' or str(breathing_freq).strip().upper() == 'SUM':
                    print(f"跳过空行或SUM行")
                    current_mode = None  # 重置共用信息
                    current_handle = None
                    continue
                
                # 如果当前行有完整信息，更新共用信息
                if breathing_mode:
                    current_mode = breathing_mode
                if occupancy_comment:
                    current_handle = occupancy_comment
                
                # 使用当前行的信息或共用信息
                mode_to_use = breathing_mode or current_mode
                handle_to_use = occupancy_comment or current_handle
                
                if not all([seat_pos, mode_to_use, breathing_freq, handle_to_use]):
                    print(f"第{row}行信息不完整，跳过")
                    continue
                
                try:
                    # 确保频率是数字
                    freq = int(float(breathing_freq))
                except (ValueError, TypeError):
                    print(f"频率格式错误，跳过此行")
                    continue
                
                # 构建对应的分类key
                handle_type = 'handle_folded' if 'folded' in str(handle_to_use).lower() else 'handle_upright'
                
                # 将座位位置转换为小写并添加前缀
                if seat_pos:
                    seat_pos = seat_pos.strip()
                    seat_mapping = {
                        'FL': 'seatleftrow1',
                        'FR': 'seatrightrow1',
                        'RL': 'seatleftrow2',
                        'RR': 'seatrightrow2',
                        'RM': 'seatmidrow2'
                    }
                    seat_type = f"seat_{seat_mapping.get(seat_pos, seat_pos.lower())}"
                
                # 处理特殊的模式名称
                mode_mapping = {
                    'Deep Sleep': 'deepsleep',
                    'Wake Up': 'wakeup',
                    'Sleep': 'sleep',
                    'Awake': 'awake'
                }
                mode_name = mode_mapping.get(mode_to_use, mode_to_use.lower())
                mode_type = f"mode_{mode_name}"
                freq_type = f"freq_{freq}"
                
                key = f"{handle_type},{freq_type},{mode_type},{seat_type}"
                print(f"构建的key: {key}")
                
                # 如果有对应的文件，写入所有文件名
                if key in self.combined_results:
                    files = self.combined_results[key]
                    if files:  # 如果有文件
                        print(f"找到匹配的文件: {files}")
                        # 直接用分号连接文件名
                        ws.cell(row=row, column=yaml_col).value = ';'.join(files)
                else:
                    print(f"未找到匹配的key: {key}")
            
            # 保存Excel文件
            wb.save(self.excel_file_path)
            print(f"Excel文件已更新并保存: {self.excel_file_path}")
            self.status_var.set("Excel文件已更新完成")
            
        except Exception as e:
            error_msg = f"更新Excel文件时出错: {str(e)}"
            print(error_msg)
            self.status_var.set(error_msg)

    def search_results(self):
        search_text = self.search_var.get().lower()
        if not search_text:
            self.clear_search()
            return
            
        # 如果没有保存原始结果，就保存一份
        if not self.original_results:
            self.original_results = self.result_text.get('1.0', tk.END)
            
        # 清空当前显示
        self.result_text.delete('1.0', tk.END)
        
        # 逐行搜索并高亮显示匹配的行
        for line in self.original_results.split('\n'):
            if search_text in line.lower():
                # 找到匹配的位置
                start = line.lower().find(search_text)
                end = start + len(search_text)
                
                # 插入匹配行，并高亮匹配的文本
                self.result_text.insert(tk.END, line[:start])
                self.result_text.insert(tk.END, line[start:end], 'highlight')
                self.result_text.insert(tk.END, line[end:] + '\n')
            
        # 配置高亮标签
        self.result_text.tag_configure('highlight', background='yellow')
        
        # 更新状态栏
        match_count = len(self.result_text.get('1.0', tk.END).strip().split('\n'))
        if match_count > 0:
            self.status_var.set(f"找到 {match_count} 个匹配结果")
        else:
            self.status_var.set("未找到匹配结果")
            
    def clear_search(self):
        # 清除搜索框
        self.search_var.set("")
        
        # 如果有原始结果，恢复显示
        if self.original_results:
            self.result_text.delete('1.0', tk.END)
            self.result_text.insert('1.0', self.original_results)
            self.original_results = ""
            
        # 更新状态栏
        self.status_var.set("")

def main():
    root = tk.Tk()
    app = YamlClassifierApp(root)
    root.mainloop()

if __name__ == "__main__":
    main() 